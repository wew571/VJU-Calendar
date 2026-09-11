# -*- coding: utf-8 -*-
"""Doc DANH SACH GIANG VIEN CO HUU tu file .xlsx cua truong.

Vi sao can: truoc do he thong phan loai co huu/thinh giang bang cach do chu
"Viet Nhat" trong o "Don vi cong tac" cua chinh file ke hoach giang day. O do la
giao vu go tay moi ky nen sai theo du kieu: bo trong (HK2 co 6 nguoi), ghi ten
truong khac nhau, va nang nhat la mot o ghi don vi cho ca nhom dong giang.

Danh sach nay la NGUON CHINH THUC: co ten trong danh sach = co huu, khong co =
thinh giang. Khong con doan tu o don vi nua.

Khuon file (sheet "List of lecturers" cua truong):

    LIST OF LECTURERS                       <- dong tieu de trang, bo qua
    No. | Full Name | Gender | Faculty | Link gg scholar | Note
    1   | PHAM DUC THO | Male | FATE, BCSE, MCSE | | ...

Chi "Full Name" la bat buoc. Cac cot khac doc duoc thi giu de hien o man Giang
vien, khong doc duoc cung khong sao.
"""
import re
import unicodedata

import openpyxl

import fate_import

# Nhan cot chap nhan duoc -> ten truong. Doc theo NHAN chu khong theo vi tri,
# cung ly le voi fate_import: file do truong gui, cot co the xe dich.
_NHAN = {
    "name": ("full name", "họ và tên", "họ tên", "tên giảng viên", "name"),
    "gender": ("gender", "giới tính"),
    "faculty": ("faculty", "khoa", "bộ môn", "đơn vị"),
    "note": ("note", "ghi chú"),
}
_SO_DONG_QUET = 12  # tim hang tieu de trong bao nhieu hang dau


def _chuan(v):
    return " ".join(str(v or "").split()).lower()


def khoa_ten(name):
    """Khoa doi chieu ten giua DANH SACH va FILE KE HOACH GIANG DAY.

    Dung lai fate_import.khoa_gv (bo hoc ham/so thu tu/dau cau) roi BO LUON DAU
    tieng Viet. Vi sao phai bo dau: hai file go boi hai nguoi khac nhau nen dat
    dau thanh khac cach - danh sach ghi "PHAN THI THANH THUY", file ke hoach ghi
    "Phan Thi Thanh Thuy" (u+dau hoi vs u+y). Giu dau thi hai ban ghi cua CUNG
    MOT NGUOI khong khop, va nguoi do bi xep nham thinh giang.

    Doi lai, hai nguoi chi khac dau se bi coi la mot. Trong pham vi mot truong
    thi hiem, va huong sai nay nhe hon: xep nham loai mot nguoi de thay ngay tren
    man Giang vien, con khong khop thi im lang.
    """
    s = unicodedata.normalize("NFD", fate_import.khoa_gv(name))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return s.replace("đ", "d")


def _tim_tieu_de(sh):
    """(chi so hang tieu de, {truong: chi so cot}) hoac (None, None)."""
    for r in range(1, min(_SO_DONG_QUET, sh.max_row) + 1):
        o = [_chuan(c) for c in next(sh.iter_rows(min_row=r, max_row=r, values_only=True))]
        cot = {}
        for field, nhan in _NHAN.items():
            for i, v in enumerate(o):
                if v in nhan:
                    cot[field] = i
                    break
        if "name" in cot:
            return r, cot
    return None, None


def read_lecturers(source):
    """Doc file -> (ket_qua, loi).

    ket_qua = {
      "sheet":  ten sheet da dung,
      "rows":   [{"key", "name", "gender", "faculty", "note", "excelRow"}],
      "skipped": [{"row", "reason"}]  - dong bo qua va vi sao,
      "duplicates": [{"name", "rows"}] - mot nguoi ghi hai lan trong danh sach,
    }
    """
    try:
        wb = openpyxl.load_workbook(source, data_only=True)
    except Exception as e:
        return None, f"Không đọc được file Excel: {e}"

    # Uu tien sheet co chu "lecturer"/"giang vien" trong ten; khong co thi quet het.
    uu_tien = [n for n in wb.sheetnames
               if "lecturer" in _chuan(n) or "giảng viên" in _chuan(n)]
    for name in uu_tien or wb.sheetnames:
        sh = wb[name]
        hang, cot = _tim_tieu_de(sh)
        if hang is not None:
            break
    else:
        wb.close()
        return None, ("Không tìm thấy hàng tiêu đề có cột \"Full Name\" (hoặc "
                      "\"Họ và tên\") trong file.")

    rows, skipped = [], []
    theo_khoa = {}
    for i, raw in enumerate(sh.iter_rows(min_row=hang + 1, values_only=True)):
        excel_row = hang + 1 + i

        def o(field):
            idx = cot.get(field)
            if idx is None or idx >= len(raw):
                return ""
            return " ".join(str(raw[idx] or "").split())

        ten = o("name")
        if not ten:
            # Dong trong hoan toan la binh thuong (cuoi bang) - khong bao.
            if any(str(v or "").strip() for v in raw):
                skipped.append({"row": excel_row, "reason": "khong_co_ten"})
            continue
        # Dong tong ket / tieu de phu lot vao giua bang.
        if re.fullmatch(r"[\d.\s]+", ten) or len(ten) < 3:
            skipped.append({"row": excel_row, "reason": "ten_khong_hop_le"})
            continue

        key = khoa_ten(ten)
        if key in theo_khoa:
            theo_khoa[key]["rows"].append(excel_row)
            continue
        muc = {"key": key, "name": ten, "gender": o("gender"),
               "faculty": o("faculty"), "note": o("note"), "excelRow": excel_row,
               "rows": [excel_row]}
        theo_khoa[key] = muc
        rows.append(muc)

    wb.close()
    trung = [{"name": m["name"], "rows": m["rows"]} for m in rows if len(m["rows"]) > 1]
    for m in rows:
        m.pop("rows", None)
    if not rows:
        return None, "File không có dòng giảng viên nào đọc được."
    return {"sheet": sh.title, "rows": rows, "skipped": skipped,
            "duplicates": trung}, None
