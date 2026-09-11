# -*- coding: utf-8 -*-
"""Ghi LUOI THOI KHOA BIEU ra .xlsx - dung bo cuc nhu tren man hinh.

KHAC HAN fate_export.py: ben do xuat BANG DU LIEU HOC PHAN (moi lop mot dong, 29
cot khuon FATE, nap lai duoc). File nay xuat cai giao vu dang NHIN THAY: luoi
tuan, hang la tiet, cot la thu, moi buoi la mot o gop cao bang so tiet cua no.
Hai thu phuc vu hai viec khac nhau - mot cai de nap lai vao he thong, mot cai de
in ra/gui di cho nguoi doc.

MODULE NAY KHONG QUYET DINH GI CA. Bo loc va bo cuc (buoi nao hien, nam cot con
thu may trong ngay, mau gi) do chinh giao dien tinh va gui sang - xem
frontend/src/adapters/xuatLuoi.js. Ly do: yeu cau la file xuat ra phai khop DUNG
man hinh, ma man hinh moi la noi biet chac no dang hien cai gi; bat backend tinh
lai la mo duong cho hai ben lech nhau ma khong ai nhin ra.

NGAY CHIA NHIEU COT: cac buoi trung gio trong cung mot ngay nam canh nhau, y het
tren luoi (Thu 4 trong du lieu that co luc can 3 cot). Tieu de ngay duoc gop
ngang qua dung so cot cua ngay do.
"""

import io

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Hai dong tieu de + mot dong ten thu, roi moi toi tiet 1.
_DONG_TIEU_DE = 1
_DONG_BO_LOC = 2
_DONG_TEN_NGAY = 4
_DONG_TIET_DAU = 5

_COT_TIET = 1          # cot A: nhan "Tiet 1", "Tiet 2"...
_RONG_COT_BUOI = 26    # do rong mot cot buoi hoc (du cho ten mon xuong dong)
_CAO_HANG_TIET = 34    # chieu cao mot hang tiet


def _mau(hex_mau):
    """'#ccfbf1' -> 'FFCCFBF1' (openpyxl doi ARGB, khong nhan dau #)."""
    s = (hex_mau or "").lstrip("#")
    return f"FF{s.upper()}" if len(s) == 6 else "FFFFFFFF"


def _vien(hex_mau):
    canh = Side(style="thin", color=_mau(hex_mau)[2:])
    return Border(left=canh, right=canh, top=canh, bottom=canh)


def _noi_dung_o(c):
    """Chu trong mot o buoi hoc - dung thu tu nhu the tren luoi (xem
    frontend/src/presentation/timetable/LessonCard.jsx, nhanh `detailed`).

    Gop ma lop / loai phong / cac dau hieu vao MOT dong: o Excel khong dat duoc
    hai cum chu ve hai phia nhu the tren man hinh, nen noi bang dau cham giua."""
    dau = []
    if c.get("daGhim"):
        dau.append("📌")
    if c.get("coVanDe"):
        dau.append("!")
    if c.get("hocChung"):
        dau.append(f"×{c['hocChung']}")

    dong1 = " ".join([c.get("maLop") or ""] + dau).strip()
    if c.get("loaiPhong"):
        dong1 = f"{dong1} · {c['loaiPhong']}" if dong1 else c["loaiPhong"]

    cuoi = f"{c.get('soTiet') or 1} tiết"
    if c.get("diaDiem"):
        cuoi = f"{cuoi} · {c['diaDiem']}"
    # Giang vien nay con day o CO SO KHAC trong cung ngay - hai co so cach rat xa,
    # khong di chuyen kip. Danh dau ngay trong o de nguoi doc file thay duoc, chu
    # khong bat ho tu do lai bang mat.
    if c.get("khacCoSo"):
        cuoi = f"{cuoi}  ⚠ còn dạy ở {c['khacCoSo']} cùng ngày"

    dong = [dong1, c.get("tenMon") or "", c.get("giangVien") or "", cuoi]
    return "\n".join(d for d in dong if d)


def _cot_dau_cua_ngay(so_cot_moi_ngay):
    """Ngay thu i bat dau o cot Excel nao (1based). Ngay trong van chiem 1 cot -
    bo mot ngay di la nguoi doc dem nham thu."""
    ra, cot = [], _COT_TIET + 1
    for n in so_cot_moi_ngay:
        ra.append(cot)
        cot += max(1, n)
    return ra, cot - 1


def build_workbook(du_lieu):
    """du_lieu: dung khuon ma frontend/src/adapters/xuatLuoi.js dung ra
    (label, moTaBoLoc, numDays, slotsPerDay, tenNgay, soCotMoiNgay, cells).
    Tra ve io.BytesIO da ghi xong, con tro ve dau, san sang cho send_file()."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Thoi khoa bieu"

    so_cot_moi_ngay = du_lieu.get("soCotMoiNgay") or []
    ten_ngay = du_lieu.get("tenNgay") or []
    slots = int(du_lieu.get("slotsPerDay") or 12)
    cot_dau, cot_cuoi = _cot_dau_cua_ngay(so_cot_moi_ngay)

    # --- Tieu de ---
    ws.cell(row=_DONG_TIEU_DE, column=1,
            value=f"THỜI KHÓA BIỂU {du_lieu.get('label') or ''}".strip()).font = Font(bold=True, size=14)
    ws.cell(row=_DONG_BO_LOC, column=1, value=du_lieu.get("moTaBoLoc") or "").font = Font(
        italic=True, size=10, color="FF666666")
    if cot_cuoi > 1:
        ws.merge_cells(start_row=_DONG_TIEU_DE, start_column=1, end_row=_DONG_TIEU_DE, end_column=cot_cuoi)
        ws.merge_cells(start_row=_DONG_BO_LOC, start_column=1, end_row=_DONG_BO_LOC, end_column=cot_cuoi)

    # --- Hang ten thu (gop ngang qua cac cot con cua ngay) ---
    nen_dau = PatternFill("solid", fgColor="FFF1F5F9")
    o = ws.cell(row=_DONG_TEN_NGAY, column=_COT_TIET, value="Tiết")
    o.font, o.fill = Font(bold=True), nen_dau
    o.alignment = Alignment(horizontal="center", vertical="center")
    for i, ten in enumerate(ten_ngay):
        d, r = cot_dau[i], max(1, so_cot_moi_ngay[i])
        o = ws.cell(row=_DONG_TEN_NGAY, column=d, value=ten)
        o.font, o.fill = Font(bold=True), nen_dau
        o.alignment = Alignment(horizontal="center", vertical="center")
        if r > 1:
            ws.merge_cells(start_row=_DONG_TEN_NGAY, start_column=d,
                           end_row=_DONG_TEN_NGAY, end_column=d + r - 1)
        for c in range(d, d + r):
            ws.cell(row=_DONG_TEN_NGAY, column=c).fill = nen_dau

    # --- Cot tiet ---
    for p in range(slots):
        o = ws.cell(row=_DONG_TIET_DAU + p, column=_COT_TIET, value=f"Tiết {p + 1}")
        o.font, o.fill = Font(bold=True, size=10), nen_dau
        o.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[_DONG_TIET_DAU + p].height = _CAO_HANG_TIET

    # --- Cac buoi hoc ---
    for c in du_lieu.get("cells") or []:
        ngay = int(c.get("day") or 0)
        if ngay >= len(cot_dau):
            continue
        tiet = int(c.get("period") or 0)
        dai = max(1, int(c.get("duration") or 1))
        cot = cot_dau[ngay] + int(c.get("col") or 0)
        dong = _DONG_TIET_DAU + tiet

        mau = c.get("mau") or {}
        o = ws.cell(row=dong, column=cot, value=_noi_dung_o(c))
        o.fill = PatternFill("solid", fgColor=_mau(mau.get("bg")))
        o.font = Font(size=9, color=_mau(mau.get("text"))[2:])
        o.border = _vien(mau.get("border"))
        # wrap_text de ten mon dai xuong dong thay vi tran sang o ben - o ben co
        # the la mot buoi khac cua cung khung gio.
        o.alignment = Alignment(vertical="top", wrap_text=True)
        if dai > 1:
            # Gop DOC dung so tiet: mot buoi 4 tiet phai cao bang 4 hang, y nhu
            # the tren luoi - do la cach doc "buoi nay chiem tu tiet may den tiet may".
            ws.merge_cells(start_row=dong, start_column=cot,
                           end_row=dong + dai - 1, end_column=cot)
            for r in range(dong, dong + dai):
                ws.cell(row=r, column=cot).border = _vien(mau.get("border"))

    # --- Kich thuoc + khoa tieu de ---
    ws.column_dimensions[get_column_letter(_COT_TIET)].width = 9
    for cot in range(_COT_TIET + 1, cot_cuoi + 1):
        ws.column_dimensions[get_column_letter(cot)].width = _RONG_COT_BUOI
    # Cuon xuong tiet 10 van con thay ten thu, cuon sang Thu 6 van con thay cot tiet.
    ws.freeze_panes = ws.cell(row=_DONG_TIET_DAU, column=_COT_TIET + 1)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
